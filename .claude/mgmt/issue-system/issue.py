#!/usr/bin/env python3
"""
issue.py - 課題管理CLI
Claude Codeから自然言語で操作するためのスクリプト

使い方:
  python3 issue.py list                          # 一覧表示
  python3 issue.py list --status 未対応          # ステータスで絞り込み
  python3 issue.py list --priority 高            # 優先度で絞り込み
  python3 issue.py list --type 機能追加          # 種別で絞り込み
  python3 issue.py list --stage IT               # ステージで絞り込み
  python3 issue.py show ISSUE-001               # 詳細表示
  python3 issue.py add                          # 対話形式で追加
  python3 issue.py add --title "..." --priority 高 --type 機能追加
  python3 issue.py update ISSUE-001 --status 対応中 --actor "Claude"
  python3 issue.py close ISSUE-001 --comment "実装完了"
  python3 issue.py verify ISSUE-001 --comment "動作確認OK"
  python3 issue.py summary                      # 集計サマリー
  python3 issue.py export                       # Excel出力

フィールド名対応（JSON内部名 → 表示名）:
  id            → ID
  title         → タイトル
  description   → 説明
  type          → 種別         [機能追加/改善/調査/設計]
  priority      → 優先度       [高/中/低]
  category      → カテゴリ
  reporter      → 報告者
  created_date  → 作成日
  stage         → 発生ステージ [URD/SRS/HLD/LLD/COD/UT/IT/ST/OPS]
  status        → ステータス
  assignee      → 担当者
  resolution    → 対応方針
  notes         → 備考
  resolved_date → 対応完了日
  verified_date → 解決確認日
  history       → 変更履歴
"""

import json
import argparse
import sys
import os
from datetime import datetime, date

DATA_FILE = os.path.join(os.path.dirname(__file__), "data", "issues.json")

VALID_STATUSES   = ["未対応", "対応中", "対応完了", "解決済", "却下"]
VALID_PRIORITIES = ["高", "中", "低"]
VALID_TYPES      = ["機能追加", "改善", "調査", "設計"]
VALID_CATEGORIES = ["解析エンジン", "merge.py", "GeoJSON出力", "prefetch", "設定", "その他"]
VALID_STAGES     = ["URD", "SRS", "HLD", "LLD", "COD", "UT", "IT", "ST", "OPS"]
# URD: ユーザー要求定義 / SRS: システム要件定義 / HLD: 基本設計 / LLD: 詳細設計
# COD: コーディング / UT: 単体テスト / IT: 結合テスト / ST: 総合テスト / OPS: 本番運用

PRIORITY_ORDER = {"高": 0, "中": 1, "低": 2}
STATUS_COLOR   = {
    "未対応":   "\033[91m",
    "対応中":   "\033[93m",
    "対応完了": "\033[96m",
    "解決済":   "\033[92m",
    "却下":     "\033[90m",
}
PRIORITY_COLOR = {
    "高": "\033[91m\033[1m", "中": "\033[93m", "低": "\033[92m",
}
RESET = "\033[0m"
BOLD  = "\033[1m"
DIM   = "\033[2m"

def load():
    with open(DATA_FILE, encoding="utf-8") as f:
        return json.load(f)

def save(data):
    data["meta"]["last_updated"] = date.today().isoformat()
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def next_id(data):
    data["meta"]["counter"] += 1
    return f"ISSUE-{data['meta']['counter']:03d}"

def colored(text, color):
    return f"{color}{text}{RESET}"

def new_issue_dict(issue_id):
    """課題の空テンプレートをXLSX列順で返す"""
    return {
        "id":            issue_id,
        "title":         None,
        "description":   None,
        "type":          None,
        "priority":      "中",
        "category":      None,
        "reporter":      None,
        "created_date":  date.today().isoformat(),
        "stage":         None,
        "status":        "未対応",
        "assignee":      None,
        "resolution":    None,
        "notes":         None,
        "resolved_date": None,
        "verified_date": None,
        "history":       [],
    }

def print_issue_row(issue):
    sid   = colored(issue["id"], BOLD)
    sstat = colored(f"{issue['status']:<5}", STATUS_COLOR.get(issue["status"], ""))
    spri  = colored(f"{issue['priority']:<3}", PRIORITY_COLOR.get(issue["priority"], ""))
    stype = (issue.get("type") or "-")[:6].ljust(6)
    title = issue["title"][:38].ljust(38)
    assign = (issue.get("assignee") or "-")[:8].ljust(8)
    print(f"  {sid}  {sstat}  {spri}  {stype}  {title}  {assign}")

def cmd_list(args):
    data   = load()
    issues = data["issues"]

    if args.status:
        issues = [i for i in issues if i["status"] == args.status]
    if args.priority:
        issues = [i for i in issues if i["priority"] == args.priority]
    if args.type:
        issues = [i for i in issues if i.get("type") == args.type]
    if args.stage:
        issues = [i for i in issues if i.get("stage") == args.stage]
    if args.assignee:
        issues = [i for i in issues if i.get("assignee") == args.assignee]
    if args.category:
        issues = [i for i in issues if i.get("category") == args.category]
    if args.open:
        issues = [i for i in issues if i["status"] not in ("解決済", "却下")]

    issues = sorted(issues, key=lambda i: (PRIORITY_ORDER.get(i["priority"], 9), i["id"]))

    if not issues:
        print("該当する課題はありません。")
        return

    print(f"\n  {BOLD}{'ID':<11}  {'状態':<7}  {'優先':<5}  {'種別':<8}  {'タイトル':<38}  {'担当者':<8}{RESET}")
    print("  " + "-" * 86)
    for i in issues:
        print_issue_row(i)
    print(f"\n  {DIM}合計: {len(issues)} 件{RESET}\n")

def cmd_show(args):
    data  = load()
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません")
        sys.exit(1)

    print(f"\n{'='*55}")
    print(f"  {BOLD}{issue['id']}  {issue['title']}{RESET}")
    print(f"{'='*55}")
    print(f"\n  {BOLD}説明:{RESET}")
    print(f"    {issue.get('description') or '-'}")
    print(f"\n  種別          : {issue.get('type') or '-'}")
    print(f"  優先度        : {colored(issue['priority'], PRIORITY_COLOR.get(issue['priority'], ''))}")
    print(f"  カテゴリ      : {issue.get('category') or '-'}")
    print(f"  報告者        : {issue.get('reporter') or '-'}")
    print(f"  作成日        : {issue.get('created_date') or '-'}")
    print(f"  発生ステージ  : {issue.get('stage') or '-'}")
    print(f"  ステータス    : {colored(issue['status'], STATUS_COLOR.get(issue['status'], ''))}")
    print(f"  担当者        : {issue.get('assignee') or '-'}")
    if issue.get("resolution"):
        print(f"\n  {BOLD}対応方針:{RESET}")
        print(f"    {issue['resolution']}")
    if issue.get("notes"):
        print(f"\n  {BOLD}備考:{RESET}")
        print(f"    {issue['notes']}")
    print(f"\n  対応完了日    : {issue.get('resolved_date') or '-'}")
    print(f"  解決確認日    : {issue.get('verified_date') or '-'}")

    history = issue.get("history", [])
    if history:
        print(f"\n  {BOLD}変更履歴:{RESET}")
        for h in history:
            print(f"    {DIM}{h['date'][:10]}{RESET}  {h['actor']}  {h['from']} → {h['to']}")
            if h.get("comment"):
                print(f"             {DIM}{h['comment']}{RESET}")
    print()

def cmd_add(args):
    data     = load()
    issue_id = next_id(data)

    if args.title:
        issue = new_issue_dict(issue_id)
        issue.update({
            "title":       args.title,
            "status":      args.status   or "未対応",
            "priority":    args.priority or "中",
            "type":        args.type,
            "assignee":    args.assignee,
            "reporter":    args.reporter,
            "stage":       args.stage,
            "category":    args.category,
            "description": args.description,
            "resolution":  args.resolution,
            "notes":       args.notes,
        })
        data["issues"].append(issue)
        save(data)
        print(f"\n  ✅ {colored(issue_id, BOLD)} を登録しました: {args.title}\n")
        return

    print(f"\n{BOLD}新規課題を登録します{RESET}\n")

    def ask(prompt, required=True, choices=None, default=None):
        hint = f" [{'/'.join(choices)}]" if choices else ""
        hint += f" (default: {default})" if default else ""
        while True:
            val = input(f"  {prompt}{hint}: ").strip()
            if not val and default:
                return default
            if not val and required:
                print("  ※ 必須項目です")
                continue
            if choices and val not in choices:
                print(f"  ※ 次のいずれかを入力してください: {', '.join(choices)}")
                continue
            return val or None

    print(f"  ID: {colored(issue_id, BOLD)}\n")

    issue = new_issue_dict(issue_id)
    issue.update({
        "title":       ask("タイトル"),
        "description": ask("説明",         required=False),
        "type":        ask("種別",         choices=VALID_TYPES,      required=False),
        "priority":    ask("優先度",       choices=VALID_PRIORITIES, default="中"),
        "category":    ask("カテゴリ",     choices=VALID_CATEGORIES, required=False),
        "reporter":    ask("報告者",       required=False),
        "stage":       ask("発生ステージ", choices=VALID_STAGES,     required=False),
        "status":      ask("ステータス",   choices=VALID_STATUSES,   default="未対応"),
        "assignee":    ask("担当者",       required=False),
        "resolution":  ask("対応方針",     required=False),
        "notes":       ask("備考",         required=False),
    })

    data["issues"].append(issue)
    save(data)
    print(f"\n  ✅ {colored(issue_id, BOLD)} を登録しました\n")

def cmd_update(args):
    data  = load()
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません")
        sys.exit(1)

    changed = []

    if args.status and args.status != issue["status"]:
        entry = {
            "date":    datetime.now().isoformat(timespec="seconds"),
            "actor":   args.actor or "不明",
            "from":    issue["status"],
            "to":      args.status,
            "comment": args.comment or "",
        }
        issue["history"].append(entry)
        old = issue["status"]
        issue["status"] = args.status
        changed.append(f"ステータス: {old} → {args.status}")

        if args.status == "対応完了" and not issue.get("resolved_date"):
            issue["resolved_date"] = date.today().isoformat()

    for field, val in [
        ("priority",   args.priority),
        ("type",       args.type),
        ("stage",      args.stage),
        ("assignee",   args.assignee),
        ("category",   args.category),
        ("description",args.description),
        ("resolution", args.resolution),
        ("notes",      args.notes),
    ]:
        if val is not None:
            changed.append(f"{field}: {issue.get(field)} → {val}")
            issue[field] = val

    if not changed:
        print("変更はありませんでした。")
        return

    save(data)
    print(f"\n  ✅ {args.id} を更新しました:")
    for c in changed:
        print(f"     • {c}")
    print()

def cmd_close(args):
    """ショートカット: ステータスを対応完了にして対応完了日を記録（Claudeが実装完了時に使う）"""
    class _A:
        pass
    a = _A()
    a.id = args.id
    a.status = "対応完了"
    a.actor = args.actor
    a.comment = args.comment
    a.priority = a.type = a.stage = a.assignee = a.category = None
    a.description = a.resolution = a.notes = None
    cmd_update(a)

def cmd_verify(args):
    """ユーザーが解決を確認：解決確認日を記録してステータスを解決済にする"""
    data  = load()
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません")
        sys.exit(1)

    today = date.today().isoformat()
    issue["verified_date"] = today
    entry = {
        "date":    datetime.now().isoformat(timespec="seconds"),
        "actor":   args.actor or "不明",
        "from":    issue["status"],
        "to":      "解決済",
        "comment": f"[解決確認] {args.comment or ''}".strip(),
    }
    issue["history"].append(entry)
    issue["status"] = "解決済"
    save(data)
    print(f"\n  ✅ {args.id} を解決済にしました（解決確認日: {today}）\n")

def cmd_summary(args):
    data   = load()
    issues = data["issues"]
    total  = len(issues)

    print(f"\n  {BOLD}{'─'*40}")
    print(f"  課題管理サマリー  (合計 {total} 件)")
    print(f"  {'─'*40}{RESET}")

    print(f"\n  {BOLD}ステータス別{RESET}")
    for s in VALID_STATUSES:
        n   = sum(1 for i in issues if i["status"] == s)
        bar = "█" * n
        col = STATUS_COLOR.get(s, "")
        print(f"  {colored(s.ljust(6), col)}  {colored(bar.ljust(20), col)}  {n}")

    open_issues = [i for i in issues if i["status"] not in ("解決済", "却下")]
    print(f"\n  {BOLD}優先度別 (未解決){RESET}")
    for pv in VALID_PRIORITIES:
        n   = sum(1 for i in open_issues if i["priority"] == pv)
        bar = "█" * n
        col = PRIORITY_COLOR.get(pv, "")
        print(f"  {colored(pv.ljust(3), col)}  {colored(bar.ljust(20), col)}  {n}")

    print(f"\n  {BOLD}種別別 (未解決){RESET}")
    for t in VALID_TYPES:
        n   = sum(1 for i in open_issues if i.get("type") == t)
        bar = "█" * n
        print(f"  {t.ljust(6)}  {bar.ljust(20)}  {n}")

    print(f"\n  {BOLD}ステージ別 (未解決){RESET}")
    for st in VALID_STAGES:
        n = sum(1 for i in open_issues if i.get("stage") == st)
        if n:
            print(f"  {st.ljust(5)}  {'█' * n}  {n}")

    resolved = sum(1 for i in issues if i["status"] == "解決済")
    rate     = resolved / total * 100 if total else 0
    print(f"\n  {BOLD}KPI{RESET}")
    print(f"  未解決件数  : {len(open_issues)}")
    print(f"  優先度「高」: {colored(str(sum(1 for i in open_issues if i['priority'] == '高')), PRIORITY_COLOR['高'])}")
    print(f"  解決率      : {resolved}/{total} ({rate:.1f}%)")
    print()

def cmd_export(args):
    """issues.jsonをExcelに書き出す（課題一覧・サマリー・収束グラフシート）"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import ScatterChart, Series, Reference
    except ImportError:
        print("openpyxlが必要です: pip install openpyxl")
        sys.exit(1)

    data   = load()
    issues = data["issues"]
    wb     = Workbook()

    thin   = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_cell(ws, row, col, value, width=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = border
        if width:
            ws.column_dimensions[get_column_letter(col)].width = width
        return c

    def data_cell(ws, row, col, value, bg="FFFFFF", bold=False):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(name="Arial", size=9, bold=bold)
        c.fill      = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border    = border
        return c

    # ── シート1: 課題一覧 ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "課題一覧"

    columns = [
        ("ID",           "id",            12),
        ("タイトル",     "title",         40),
        ("説明",         "description",   44),
        ("種別",         "type",          12),
        ("優先度",       "priority",      10),
        ("カテゴリ",     "category",      16),
        ("報告者",       "reporter",      14),
        ("作成日",       "created_date",  14),
        ("発生ステージ", "stage",         16),
        ("ステータス",   "status",        12),
        ("担当者",       "assignee",      14),
        ("対応方針",     "resolution",    30),
        ("備考",         "notes",         30),
        ("対応完了日",   "resolved_date", 14),
        ("解決確認日",   "verified_date", 14),
    ]

    ws1.row_dimensions[1].height = 28
    for i, (label, _, width) in enumerate(columns, 1):
        hdr_cell(ws1, 1, i, label, width=width)

    priority_bg = {"高": "FFE0E0", "中": "FFFBE0", "低": "E8F5E9"}
    pri_col = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "priority")

    for ri, issue in enumerate(issues, 2):
        bg = "F0F0F0" if issue["status"] in ("解決済", "却下") else ("EBF3FB" if ri % 2 == 0 else "FFFFFF")
        for ci, (_, field, _) in enumerate(columns, 1):
            val = issue.get(field) or ""
            cell_bg = priority_bg.get(val, bg) if ci == pri_col else bg
            data_cell(ws1, ri, ci, val, bg=cell_bg, bold=(ci == pri_col and val in priority_bg))
        ws1.row_dimensions[ri].height = 40

    ws1.freeze_panes = "B2"

    from openpyxl.worksheet.datavalidation import DataValidation
    last_row = len(issues) + 1

    status_col   = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "status")
    priority_col = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "priority")
    type_col     = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "type")
    stage_col    = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "stage")

    def make_dv(choices, col_letter, last_row):
        dv = DataValidation(
            type="list",
            formula1=f'"{",".join(choices)}"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.sqref = f"{col_letter}2:{col_letter}{last_row}"
        return dv

    ws1.add_data_validation(make_dv(VALID_STATUSES,   get_column_letter(status_col),   last_row))
    ws1.add_data_validation(make_dv(VALID_PRIORITIES, get_column_letter(priority_col), last_row))
    ws1.add_data_validation(make_dv(VALID_TYPES,      get_column_letter(type_col),     last_row))
    ws1.add_data_validation(make_dv(VALID_STAGES,     get_column_letter(stage_col),    last_row))

    # ── シート2: サマリー ──────────────────────────────────────
    ws2 = wb.create_sheet("サマリー")
    open_issues = [i for i in issues if i["status"] not in ("解決済", "却下")]
    resolved    = sum(1 for i in issues if i["status"] == "解決済")
    total       = len(issues)

    def section_header(ws, row, col, text, span=2):
        c = ws.cell(row=row, column=col, value=text)
        c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor="2E5D9E")
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border    = border
        for dc in range(1, span):
            nc = ws.cell(row=row, column=col + dc)
            nc.fill   = PatternFill("solid", fgColor="2E5D9E")
            nc.border = border

    def summary_row(ws, row, label, value, label_bg="F5F5F5", value_bg="FFFFFF"):
        lc = ws.cell(row=row, column=1, value=label)
        lc.font      = Font(name="Arial", size=9, bold=True)
        lc.fill      = PatternFill("solid", fgColor=label_bg)
        lc.alignment = Alignment(horizontal="left", vertical="center")
        lc.border    = border
        vc = ws.cell(row=row, column=2, value=value)
        vc.font      = Font(name="Arial", size=9)
        vc.fill      = PatternFill("solid", fgColor=value_bg)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border    = border

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 12

    row = 1
    section_header(ws2, row, 1, "KPI"); row += 1
    summary_row(ws2, row, "合計件数",   total);           row += 1
    summary_row(ws2, row, "未解決件数", len(open_issues)); row += 1
    summary_row(ws2, row, "解決済件数", resolved);        row += 1
    summary_row(ws2, row, "解決率", f"{resolved/total*100:.1f}%" if total else "-%"); row += 2

    section_header(ws2, row, 1, "ステータス別"); row += 1
    for s in VALID_STATUSES:
        summary_row(ws2, row, s, sum(1 for i in issues if i["status"] == s)); row += 1
    row += 1

    section_header(ws2, row, 1, "優先度別（未解決）"); row += 1
    for pv in VALID_PRIORITIES:
        summary_row(ws2, row, pv, sum(1 for i in open_issues if i["priority"] == pv)); row += 1
    row += 1

    section_header(ws2, row, 1, "種別別（未解決）"); row += 1
    for t in VALID_TYPES:
        n = sum(1 for i in open_issues if i.get("type") == t)
        if n:
            summary_row(ws2, row, t, n); row += 1
    row += 1

    section_header(ws2, row, 1, "カテゴリ別（未解決）"); row += 1
    for cat in VALID_CATEGORIES:
        n = sum(1 for i in open_issues if i.get("category") == cat)
        if n:
            summary_row(ws2, row, cat, n); row += 1
    row += 1

    section_header(ws2, row, 1, "ステージ別（未解決）"); row += 1
    for st in VALID_STAGES:
        n = sum(1 for i in open_issues if i.get("stage") == st)
        if n:
            summary_row(ws2, row, st, n); row += 1

    # ── シート3: 収束グラフ ──────────────────────────────────────
    ws3 = wb.create_sheet("収束グラフ")

    all_created = [i["created_date"] for i in issues if i.get("created_date")]
    if all_created:
        from datetime import timedelta as td
        start = date.fromisoformat(min(all_created))
        end   = date.today()
        dates = []
        d = start
        while d <= end:
            dates.append(d)
            d += td(days=1)

        for ci, label in enumerate(["日付", "課題累計", "対応累計", "確認累計"], 1):
            c = ws3.cell(row=1, column=ci, value=label)
            c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            c.fill      = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = border
        ws3.column_dimensions["A"].width = 14
        ws3.column_dimensions["B"].width = 12
        ws3.column_dimensions["C"].width = 12
        ws3.column_dimensions["D"].width = 12

        def to_serial(d):
            return (d - date(1899, 12, 30)).days

        for ri, d in enumerate(dates, 2):
            ds = d.isoformat()
            created_n  = sum(1 for i in issues if i.get("created_date")  and i["created_date"]  <= ds)
            resolved_n = sum(1 for i in issues if i.get("resolved_date") and i["resolved_date"] <= ds)
            verified_n = sum(1 for i in issues if i.get("verified_date") and i["verified_date"] <= ds)
            c = ws3.cell(row=ri, column=1, value=to_serial(d))
            c.number_format = "YYYY-MM-DD"
            ws3.cell(row=ri, column=2, value=created_n)
            ws3.cell(row=ri, column=3, value=resolved_n)
            ws3.cell(row=ri, column=4, value=verified_n)

        chart = ScatterChart()
        chart.title  = "課題収束曲線"
        chart.style  = 10
        chart.width  = 20
        chart.height = 14
        chart.y_axis.title  = "累計数"
        chart.x_axis.title  = "日付"
        chart.y_axis.numFmt = "0"
        chart.x_axis.numFmt = "YYYY-MM-DD"
        chart.x_axis.scaling.min = to_serial(start)
        chart.x_axis.scaling.max = to_serial(end)
        chart.y_axis.scaling.min = 0

        xvals = Reference(ws3, min_col=1, min_row=2, max_row=len(dates) + 1)
        series_defs = [
            ("課題累計", 2, "FF0000"),
            ("対応累計", 3, "4472C4"),
            ("確認累計", 4, "70AD47"),
        ]
        for label, col, color in series_defs:
            yvals = Reference(ws3, min_col=col, min_row=2, max_row=len(dates) + 1)
            s = Series(yvals, xvals, title=label)
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width     = 20000
            s.marker.symbol = "circle"
            s.marker.size   = 5
            chart.series.append(s)

        ws3.add_chart(chart, "F2")

    out = args.output or os.path.join(os.path.dirname(__file__), "reports", "issues_export.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"\n  ✅ エクスポート完了: {out}\n")

def main():
    p = argparse.ArgumentParser(description="課題管理CLI")
    sub = p.add_subparsers(dest="cmd")

    # list
    pl = sub.add_parser("list", help="一覧表示")
    pl.add_argument("--status",   help="ステータスで絞り込み")
    pl.add_argument("--priority", help="優先度で絞り込み")
    pl.add_argument("--type",     help="種別で絞り込み")
    pl.add_argument("--stage",    help="ステージで絞り込み")
    pl.add_argument("--assignee", help="担当者で絞り込み")
    pl.add_argument("--category", help="カテゴリで絞り込み")
    pl.add_argument("--open",     action="store_true", help="未解決のみ")

    # show
    ps = sub.add_parser("show", help="詳細表示")
    ps.add_argument("id", help="課題ID (例: ISSUE-001)")

    # add
    pa = sub.add_parser("add", help="新規課題を登録 (--title 指定で非対話、省略で対話形式)")
    pa.add_argument("--title")
    pa.add_argument("--status",      choices=VALID_STATUSES)
    pa.add_argument("--priority",    choices=VALID_PRIORITIES)
    pa.add_argument("--type",        choices=VALID_TYPES)
    pa.add_argument("--stage",       choices=VALID_STAGES)
    pa.add_argument("--category",    choices=VALID_CATEGORIES)
    pa.add_argument("--assignee")
    pa.add_argument("--reporter")
    pa.add_argument("--description")
    pa.add_argument("--resolution")
    pa.add_argument("--notes")

    # update
    pu = sub.add_parser("update", help="課題を更新")
    pu.add_argument("id")
    pu.add_argument("--status",    choices=VALID_STATUSES)
    pu.add_argument("--priority",  choices=VALID_PRIORITIES)
    pu.add_argument("--type",      choices=VALID_TYPES)
    pu.add_argument("--stage",     choices=VALID_STAGES)
    pu.add_argument("--category",  choices=VALID_CATEGORIES)
    pu.add_argument("--assignee")
    pu.add_argument("--description")
    pu.add_argument("--resolution")
    pu.add_argument("--notes")
    pu.add_argument("--actor",   default="不明", help="変更者名")
    pu.add_argument("--comment", help="変更コメント")

    # close
    pc = sub.add_parser("close", help="対応完了にして対応完了日を記録")
    pc.add_argument("id")
    pc.add_argument("--actor",   default="不明")
    pc.add_argument("--comment", help="完了コメント")

    # verify
    pv = sub.add_parser("verify", help="解決確認日を記録")
    pv.add_argument("id")
    pv.add_argument("--actor",   default="不明")
    pv.add_argument("--comment", help="確認コメント")

    # summary
    sub.add_parser("summary", help="集計サマリーを表示")

    # export
    pe = sub.add_parser("export", help="Excelファイルに出力")
    pe.add_argument("--output", help="出力ファイルパス")

    args = p.parse_args()
    if not args.cmd:
        p.print_help()
        return

    dispatch = {
        "list": cmd_list, "show": cmd_show, "add": cmd_add,
        "update": cmd_update, "close": cmd_close, "verify": cmd_verify,
        "summary": cmd_summary, "export": cmd_export,
    }
    dispatch[args.cmd](args)

if __name__ == "__main__":
    main()
