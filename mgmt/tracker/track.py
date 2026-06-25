#!/usr/bin/env python3
"""
track.py - バグ・課題 統合管理CLI

使い方:
  python3 track.py bug   <command> [options]   # バグ管理
  python3 track.py issue <command> [options]   # 課題管理

コマンド共通:
  list    [--status/--open ...]   一覧表示
  show    <ID>                    詳細表示
  add     [--title ...]           新規登録
  update  <ID> [--status ...]     更新
  close   <ID>                    対応完了にする（Claude が使う）
  verify  <ID>                    解決確認（ユーザーが使う）
  summary                         サマリー表示
  export                          Excel出力

例:
  python3 track.py bug list --open
  python3 track.py bug add --title "..." --severity 高 --stage COD
  python3 track.py bug close BUG-001 --actor "Claude" --comment "修正完了"
  python3 track.py issue list --open
  python3 track.py issue add --title "..." --priority 高 --type 改善
  python3 track.py issue close ISSUE-001 --actor "Claude" --comment "実装完了"
"""

import argparse
import json
import os
import re
import sys
from datetime import date, datetime

BASE_DIR      = os.path.dirname(__file__)
BUG_DATA      = os.path.join(BASE_DIR, "data", "bugs.json")
ISSUE_DATA    = os.path.join(BASE_DIR, "data", "issues.json")
BUG_REPORT    = os.path.join(BASE_DIR, "reports", "bugs_export.xlsx")
ISSUE_REPORT  = os.path.join(BASE_DIR, "reports", "issues_export.xlsx")

VALID_STATUSES   = ["未対応", "対応中", "対応完了", "解決済", "却下"]
VALID_SEVERITIES = ["高", "中", "低"]
VALID_PRIORITIES = ["高", "中", "低"]
VALID_TYPES      = ["改善", "調査", "設計"]  # 「機能追加」は廃止（FR確定済み実装は todo.md へ）。再追加禁止
VALID_CATEGORIES = ["解析エンジン", "merge.py", "GeoJSON出力", "prefetch", "設定", "その他"]
VALID_STAGES     = ["URD", "SRS", "HLD", "LLD", "COD", "UT", "IT", "ST", "OPS"]

STATUS_COLOR = {
    "未対応":   "\033[91m",
    "対応中":   "\033[93m",
    "対応完了": "\033[96m",
    "解決済":   "\033[92m",
    "却下":     "\033[90m",
}
SEVERITY_COLOR = {"高": "\033[91m\033[1m", "中": "\033[93m", "低": "\033[92m"}
PRIORITY_COLOR = {"高": "\033[91m\033[1m", "中": "\033[93m", "低": "\033[92m"}
RESET = "\033[0m"
BOLD  = "\033[1m"
DIM   = "\033[2m"


# ── 共通ヘルパー ──────────────────────────────────────────────────────────

def colored(text, color):
    return f"{color}{text}{RESET}"

def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)

def save(path, data):
    data["meta"]["last_updated"] = date.today().isoformat()
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def next_id(data, prefix):
    data["meta"]["counter"] += 1
    return f"{prefix}-{data['meta']['counter']:03d}"

def normalize_id(raw, kind):
    prefix = {"bug": "BUG", "issue": "ISSUE"}.get(kind)
    if prefix is None or raw is None:
        return raw
    m = re.search(r'(\d+)\s*$', str(raw).strip())
    if not m:
        return raw
    return f"{prefix}-{int(m.group(1)):03d}"

def ask_interactive(prompt, required=True, choices=None, default=None, guide=None):
    if guide:
        print(f"  {DIM}{guide}{RESET}")
    if choices:
        print(f"  {prompt}:")
        for i, c in enumerate(choices, 1):
            marker = " ←default" if c == default else ""
            print(f"    {i}) {c}{marker}")
        hint = f"番号または文字列{f' (default: {default})' if default else ''}"
        while True:
            val = input(f"  選択 [{hint}]: ").strip()
            if not val and default:
                return default
            if not val and required:
                print("  ※ 必須項目です")
                continue
            if not val:
                return None
            if val.isdigit() and 1 <= int(val) <= len(choices):
                return choices[int(val) - 1]
            if val in choices:
                return val
            print(f"  ※ 番号（1〜{len(choices)}）または {', '.join(choices)} を入力してください")
    else:
        hint = f" (default: {default})" if default else ""
        while True:
            val = input(f"  {prompt}{hint}: ").strip()
            if not val and default:
                return default
            if not val and required:
                print("  ※ 必須項目です")
                continue
            return val or None

def current_session_id():
    return os.environ.get("CLAUDE_CODE_SESSION_ID", "")

def append_history(item, actor, from_s, to_s, comment=""):
    item["history"].append({
        "date":       datetime.now().isoformat(timespec="seconds"),
        "actor":      actor or "不明",
        "from":       from_s,
        "to":         to_s,
        "comment":    comment or "",
        "session_id": current_session_id(),
    })


# ── Excel エクスポート共通ヘルパー ────────────────────────────────────────

def make_excel_helpers():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin   = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_cell(ws, row, col, value, width=None):
        from openpyxl.utils import get_column_letter
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = border
        if width:
            ws.column_dimensions[get_column_letter(col)].width = width
        return c

    def data_cell(ws, row, col, value, bg="FFFFFF", bold=False):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(name="Arial", size=9, bold=bold)
        c.fill      = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border    = border
        return c

    def section_header(ws, row, col, text, span=2):
        from openpyxl.styles import Alignment, Font, PatternFill
        c = ws.cell(row=row, column=col, value=text)
        c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor="2E5D9E")
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border    = border
        for dc in range(1, span):
            nc = ws.cell(row=row, column=col + dc)
            nc.fill   = PatternFill("solid", fgColor="2E5D9E")
            nc.border = border

    def summary_row(ws, row, label, value):
        from openpyxl.styles import Alignment, Font, PatternFill
        lc = ws.cell(row=row, column=1, value=label)
        lc.font      = Font(name="Arial", size=9, bold=True)
        lc.fill      = PatternFill("solid", fgColor="F5F5F5")
        lc.alignment = Alignment(horizontal="left", vertical="center")
        lc.border    = border
        vc = ws.cell(row=row, column=2, value=value)
        vc.font      = Font(name="Arial", size=9)
        vc.fill      = PatternFill("solid", fgColor="FFFFFF")
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border    = border

    def make_dv(choices, col_letter, last_row):
        from openpyxl.worksheet.datavalidation import DataValidation
        dv = DataValidation(
            type="list",
            formula1=f'"{",".join(choices)}"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.sqref = f"{col_letter}2:{col_letter}{last_row}"
        return dv

    def to_serial(d):
        return (d - date(1899, 12, 30)).days

    return border, hdr_cell, data_cell, section_header, summary_row, make_dv, to_serial


def write_convergence_sheet(ws, items, date_field, hdr_cell, border, to_serial):
    """収束グラフ用シートを書き込む（Bug/Issue 共通）"""
    from datetime import timedelta as td

    from openpyxl.chart import LineChart, Reference, Series
    from openpyxl.chart.data_source import AxDataSource, StrRef
    from openpyxl.styles import Alignment, Font, PatternFill

    all_dates = [x[date_field] for x in items if x.get(date_field)]
    if not all_dates:
        return

    start = date.fromisoformat(min(all_dates))
    end   = date.today()
    dates = []
    d = start
    while d <= end:
        dates.append(d)
        d += td(days=1)

    labels = ["日付", "累計登録", "対応累計", "確認累計"]
    for ci, label in enumerate(labels, 1):
        c = ws.cell(row=1, column=ci, value=label)
        c.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = border
    for col_letter, width in [("A", 14), ("B", 12), ("C", 12), ("D", 12)]:
        ws.column_dimensions[col_letter].width = width

    for ri, d in enumerate(dates, 2):
        ds = d.isoformat()
        cnt_reg  = sum(1 for x in items if x.get(date_field)     and x[date_field]     <= ds)
        cnt_res  = sum(1 for x in items if x.get("resolved_date") and x["resolved_date"] <= ds)
        cnt_ver  = sum(1 for x in items if x.get("verified_date") and x["verified_date"] <= ds)
        ws.cell(row=ri, column=1, value=d.isoformat())
        ws.cell(row=ri, column=2, value=cnt_reg)
        ws.cell(row=ri, column=3, value=cnt_res)
        ws.cell(row=ri, column=4, value=cnt_ver)

    chart = LineChart()
    chart.title  = ws.title
    chart.style  = 10
    chart.width  = 20
    chart.height = 14
    chart.y_axis.title   = "累計数"
    chart.x_axis.title   = "日付"
    chart.y_axis.numFmt  = "0"
    chart.y_axis.scaling.min = 0

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(dates) + 1)
    cats_src = AxDataSource(strRef=StrRef(f=str(cats)))
    for _label, col, color in [("累計登録", 2, "FF0000"), ("対応累計", 3, "4472C4"), ("確認累計", 4, "70AD47")]:
        yvals = Reference(ws, min_col=col, min_row=1, max_row=len(dates) + 1)
        s = Series(yvals, title_from_data=True)
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width     = 20000
        s.cat = cats_src
        chart.series.append(s)

    ws.add_chart(chart, "F2")


# ═══════════════════════════════════════════════════════════════════════════
# BUG コマンド
# ═══════════════════════════════════════════════════════════════════════════

def new_bug(bug_id):
    return {
        "id":                 bug_id,
        "title":              None,
        "description":        None,
        "reporter":           None,
        "found_date":         date.today().isoformat(),
        "found_stage":        None,
        "category":           None,
        "found_in":           None,
        "repro":              None,
        "severity":           "中",
        "status":             "未対応",
        "assignee":           None,
        "stage":              None,
        "cause":              None,
        "resolution":         None,
        "notes":              None,
        "resolved_date":      None,
        "verified_date":      None,
        "created_session_id": current_session_id(),
        "history":            [],
    }

def print_bug_row(bug):
    sid   = colored(bug["id"], BOLD)
    sstat = colored(f"{bug['status']:<5}", STATUS_COLOR.get(bug["status"], ""))
    ssev  = colored(f"{bug['severity']:<3}", SEVERITY_COLOR.get(bug["severity"], ""))
    title = bug["title"][:42].ljust(42)
    asgn  = (bug.get("assignee") or "-")[:8].ljust(8)
    print(f"  {sid}  {sstat}  {ssev}  {title}  {asgn}")

def bug_list(args):
    data = load(BUG_DATA)
    bugs = data["bugs"]
    if args.status:   bugs = [b for b in bugs if b["status"]              == args.status]
    if args.severity: bugs = [b for b in bugs if b["severity"]            == args.severity]
    if args.assignee: bugs = [b for b in bugs if b.get("assignee")        == args.assignee]
    if args.category: bugs = [b for b in bugs if b.get("category")        == args.category]
    if args.open:     bugs = [b for b in bugs if b["status"] not in ("解決済", "却下")]
    bugs = sorted(bugs, key=lambda b: ({"高":0,"中":1,"低":2}.get(b["severity"], 9), b["id"]))
    if not bugs:
        print("該当するバグはありません。"); return
    print(f"\n  {BOLD}{'ID':<9}  {'状態':<7}  {'重大度':<5}  {'タイトル':<42}  {'担当者':<8}{RESET}")
    print("  " + "-" * 80)
    for b in bugs:
        print_bug_row(b)
    print(f"\n  {DIM}合計: {len(bugs)} 件{RESET}\n")

def bug_show(args):
    data = load(BUG_DATA)
    bug  = next((b for b in data["bugs"] if b["id"] == args.id), None)
    if not bug:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    print(f"\n{'='*55}")
    print(f"  {BOLD}{bug['id']}  {bug['title']}{RESET}")
    print(f"{'='*55}")
    print(f"\n  {BOLD}説明:{RESET}\n    {bug.get('description') or '-'}")
    print(f"\n  報告者        : {bug.get('reporter') or '-'}")
    print(f"  発見日        : {bug.get('found_date') or '-'}")
    print(f"  発生ステージ  : {bug.get('found_stage') or '-'}")
    print(f"  カテゴリ      : {bug.get('category') or '-'}")
    print(f"  発生プログラム: {bug.get('found_in') or '-'}")
    print(f"  再現性        : {bug.get('repro') or '-'}")
    print(f"  重大度        : {colored(bug['severity'], SEVERITY_COLOR.get(bug['severity'], ''))}")
    print(f"  ステータス    : {colored(bug['status'], STATUS_COLOR.get(bug['status'], ''))}")
    print(f"  担当者        : {bug.get('assignee') or '-'}")
    print(f"  発生源ステージ: {bug.get('stage') or '-'}")
    print(f"  原因プログラム: {bug.get('cause') or '-'}")
    if bug.get("resolution"):
        print(f"\n  {BOLD}対応方針:{RESET}\n    {bug['resolution']}")
    if bug.get("notes"):
        print(f"\n  {BOLD}備考:{RESET}\n    {bug['notes']}")
    print(f"\n  対応完了日    : {bug.get('resolved_date') or '-'}")
    print(f"  解決確認日    : {bug.get('verified_date') or '-'}")
    csid = bug.get("created_session_id") or "-"
    print(f"  登録セッション: {csid[:8] if csid != '-' else '-'}")
    history = bug.get("history", [])
    if history:
        print(f"\n  {BOLD}変更履歴:{RESET}")
        for h in history:
            sid = h.get("session_id") or ""
            sid_str = f"  [{sid[:8]}]" if sid else ""
            print(f"    {DIM}{h['date'][:10]}{RESET}  {h['actor']}  {h['from']} → {h['to']}{sid_str}")
            if h.get("comment"):
                print(f"             {DIM}{h['comment']}{RESET}")
    print()

def bug_add(args):
    data   = load(BUG_DATA)
    bug_id = next_id(data, "BUG")
    if args.title:
        bug = new_bug(bug_id)
        bug.update({
            "title":       args.title,
            "status":      args.status     or "未対応",
            "severity":    args.severity   or "中",
            "assignee":    args.assignee,
            "reporter":    args.reporter,
            "found_stage": args.found_stage,
            "stage":       args.stage,
            "category":    args.category,
            "found_in":    args.found_in,
            "cause":       args.cause,
            "repro":       args.repro,
            "description": args.description,
            "resolution":  args.resolution,
            "notes":       args.notes,
        })
        data["bugs"].append(bug)
        save(BUG_DATA, data)
        print(f"\n  ✅ {colored(bug_id, BOLD)} を登録しました: {args.title}\n")
        return
    print(f"\n{BOLD}新規バグを登録します{RESET}\n")
    print(f"  ID: {colored(bug_id, BOLD)}\n")
    bug = new_bug(bug_id)
    bug.update({
        "title":       ask_interactive("タイトル",
                           guide="バグの主題を一行で。例: merge.py: 削除候補の件数が0件になる"),
        "description": ask_interactive("説明",
                           required=False,
                           guide="期待する動作・実際の動作・再現手順を記述"),
        "reporter":    ask_interactive("報告者",
                           required=False,
                           guide="人名 または モデル名（Opus / Sonnet 等）。Claude が登録する場合はモデル名"),
        "found_stage": ask_interactive("発生ステージ",
                           choices=VALID_STAGES, required=False,
                           guide="バグが発見されたフェーズ（テスト中なら UT/IT/ST、実装中なら COD 等）"),
        "category":    ask_interactive("カテゴリ",
                           choices=VALID_CATEGORIES, required=False,
                           guide="どのモジュール・機能に関係するか"),
        "found_in":    ask_interactive("発生プログラム",
                           required=False,
                           guide="バグが現れるファイル。例: scripts/merge.py"),
        "repro":       ask_interactive("再現性",
                           required=False,
                           guide="毎回 / 条件付き / たまに"),
        "severity":    ask_interactive("重大度",
                           choices=VALID_SEVERITIES, default="中",
                           guide="高=結果が信用できない/完走しない、中=効率・使い勝手に支障、低=ログ・命名等の品質"),
        "status":      ask_interactive("ステータス",
                           choices=VALID_STATUSES, default="未対応"),
        "assignee":    ask_interactive("担当者",         required=False),
        "stage":       ask_interactive("発生源ステージ",
                           choices=VALID_STAGES, required=False,
                           guide="バグの原因が埋め込まれたフェーズ（仕様ミスなら SRS、実装ミスなら COD 等）"),
        "cause":       ask_interactive("原因プログラム",
                           required=False,
                           guide="バグの根本原因があるファイル。例: src/analyze.c"),
        "resolution":  ask_interactive("対応方針",
                           required=False,
                           guide="修正方針。例: analyze.c の境界条件チェックを追加する"),
        "notes":       ask_interactive("備考",           required=False),
    })
    data["bugs"].append(bug)
    save(BUG_DATA, data)
    print(f"\n  ✅ {colored(bug_id, BOLD)} を登録しました\n")

def bug_update(args):
    data = load(BUG_DATA)
    bug  = next((b for b in data["bugs"] if b["id"] == args.id), None)
    if not bug:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    changed = []
    if args.status and args.status != bug["status"]:
        append_history(bug, args.actor, bug["status"], args.status, args.comment)
        old = bug["status"]
        bug["status"] = args.status
        changed.append(f"ステータス: {old} → {args.status}")
        if args.status == "対応完了" and not bug.get("resolved_date"):
            bug["resolved_date"] = date.today().isoformat()
        if args.status == "却下" and not bug.get("verified_date"):
            bug["verified_date"] = date.today().isoformat()
    for field, val in [
        ("severity",    args.severity),
        ("assignee",    args.assignee),
        ("found_stage", args.found_stage),
        ("stage",       args.stage),
        ("category",    args.category),
        ("found_in",    args.found_in),
        ("cause",       args.cause),
        ("repro",       args.repro),
        ("description", args.description),
        ("resolution",  args.resolution),
        ("notes",       args.notes),
    ]:
        if val is not None:
            changed.append(f"{field}: {bug.get(field)} → {val}")
            bug[field] = val
    if not changed:
        print("変更はありませんでした。"); return
    save(BUG_DATA, data)
    print(f"\n  ✅ {args.id} を更新しました:")
    for c in changed:
        print(f"     • {c}")
    print()

def bug_close(args):
    class _A: pass
    a = _A()
    a.id = args.id; a.status = "対応完了"; a.actor = args.actor; a.comment = args.comment
    a.severity = a.assignee = a.found_stage = a.stage = a.category = None
    a.found_in = a.cause = a.repro = a.description = a.resolution = a.notes = None
    bug_update(a)

def bug_verify(args):
    data = load(BUG_DATA)
    bug  = next((b for b in data["bugs"] if b["id"] == args.id), None)
    if not bug:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    today = date.today().isoformat()
    bug["verified_date"] = today
    append_history(bug, args.actor, bug["status"], "解決済", f"[解決確認] {args.comment or ''}".strip())
    bug["status"] = "解決済"
    save(BUG_DATA, data)
    print(f"\n  ✅ {args.id} を解決済にしました（解決確認日: {today}）\n")

def bug_summary(args):
    data  = load(BUG_DATA)
    bugs  = data["bugs"]
    total = len(bugs)
    print(f"\n  {BOLD}{'─'*40}\n  バグ管理サマリー  (合計 {total} 件)\n  {'─'*40}{RESET}")
    print(f"\n  {BOLD}ステータス別{RESET}")
    for s in VALID_STATUSES:
        n = sum(1 for b in bugs if b["status"] == s)
        print(f"  {colored(s.ljust(6), STATUS_COLOR.get(s,''))}  {colored(('█'*n).ljust(20), STATUS_COLOR.get(s,''))}  {n}")
    open_bugs = [b for b in bugs if b["status"] not in ("解決済", "却下")]
    print(f"\n  {BOLD}重大度別 (未解決){RESET}")
    for sv in VALID_SEVERITIES:
        n = sum(1 for b in open_bugs if b["severity"] == sv)
        print(f"  {colored(sv.ljust(3), SEVERITY_COLOR.get(sv,''))}  {colored(('█'*n).ljust(20), SEVERITY_COLOR.get(sv,''))}  {n}")
    resolved = sum(1 for b in bugs if b["status"] == "解決済")
    print(f"\n  {BOLD}KPI{RESET}")
    print(f"  未解決件数  : {len(open_bugs)}")
    print(f"  重大度「高」: {colored(str(sum(1 for b in open_bugs if b['severity']=='高')), SEVERITY_COLOR['高'])}")
    print(f"  解決率      : {resolved}/{total} ({resolved/total*100:.1f}%)" if total else "  解決率      : 0/0 (-)")
    print()

def bug_export(args):
    out = args.output or BUG_REPORT
    if args.if_changed and os.path.exists(out):
        if os.path.getmtime(BUG_DATA) <= os.path.getmtime(out):
            print(f"  変更なし、{out} の再生成をスキップ")
            return

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxlが必要です: pip install openpyxl"); sys.exit(1)

    data = load(BUG_DATA)
    bugs = data["bugs"]
    wb   = Workbook()
    border, hdr_cell, data_cell, section_header, summary_row, make_dv, to_serial = make_excel_helpers()

    ws1 = wb.active
    ws1.title = "バグ一覧"
    columns = [
        ("ID",            "id",            10),
        ("タイトル",      "title",         40),
        ("説明",          "description",   44),
        ("報告者",        "reporter",      14),
        ("発見日",        "found_date",    14),
        ("発生ステージ",  "found_stage",   16),
        ("カテゴリ",      "category",      16),
        ("発生プログラム","found_in",      20),
        ("再現性",        "repro",         10),
        ("重大度",        "severity",      10),
        ("ステータス",    "status",        12),
        ("担当者",        "assignee",      14),
        ("発生源ステージ","stage",         16),
        ("原因プログラム","cause",         20),
        ("対応方針",      "resolution",    30),
        ("備考",          "notes",         30),
        ("対応完了日",    "resolved_date", 14),
        ("解決確認日",    "verified_date", 14),
    ]
    ws1.row_dimensions[1].height = 28
    for i, (label, _, width) in enumerate(columns, 1):
        hdr_cell(ws1, 1, i, label, width=width)
    severity_bg = {"高": "FFE0E0", "中": "FFFBE0", "低": "E8F5E9"}
    sev_col = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "severity")
    for ri, bug in enumerate(bugs, 2):
        bg = "F0F0F0" if bug["status"] in ("解決済", "却下") else ("EBF3FB" if ri % 2 == 0 else "FFFFFF")
        for ci, (_, field, _) in enumerate(columns, 1):
            val = bug.get(field) or ""
            cell_bg = severity_bg.get(val, bg) if ci == sev_col else bg
            data_cell(ws1, ri, ci, val, bg=cell_bg, bold=(ci == sev_col and val in severity_bg))
        ws1.row_dimensions[ri].height = 40
    ws1.freeze_panes = "B2"
    last_row = len(bugs) + 1
    for choices, field in [(VALID_STATUSES, "status"), (VALID_SEVERITIES, "severity"),
                           (VALID_STAGES, "found_stage"), (VALID_STAGES, "stage")]:
        col_idx = next(i for i, (_, f, _) in enumerate(columns, 1) if f == field)
        ws1.add_data_validation(make_dv(choices, get_column_letter(col_idx), last_row))

    ws2 = wb.create_sheet("サマリー")
    open_bugs = [b for b in bugs if b["status"] not in ("解決済", "却下")]
    resolved  = sum(1 for b in bugs if b["status"] == "解決済")
    total     = len(bugs)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 12
    row = 1
    section_header(ws2, row, 1, "KPI"); row += 1
    summary_row(ws2, row, "合計件数",   total);         row += 1
    summary_row(ws2, row, "未解決件数", len(open_bugs)); row += 1
    summary_row(ws2, row, "解決済件数", resolved);      row += 1
    summary_row(ws2, row, "解決率", f"{resolved/total*100:.1f}%" if total else "-%"); row += 2
    section_header(ws2, row, 1, "ステータス別"); row += 1
    for s in VALID_STATUSES:
        summary_row(ws2, row, s, sum(1 for b in bugs if b["status"] == s)); row += 1
    row += 1
    section_header(ws2, row, 1, "重大度別（未解決）"); row += 1
    for sv in VALID_SEVERITIES:
        summary_row(ws2, row, sv, sum(1 for b in open_bugs if b["severity"] == sv)); row += 1
    row += 1
    section_header(ws2, row, 1, "カテゴリ別（未解決）"); row += 1
    for cat in VALID_CATEGORIES:
        n = sum(1 for b in open_bugs if b.get("category") == cat)
        if n:
            summary_row(ws2, row, cat, n); row += 1
    row += 1
    section_header(ws2, row, 1, "発生源ステージ別（未解決）"); row += 1
    for st in VALID_STAGES:
        n = sum(1 for b in open_bugs if b.get("stage") == st)
        if n:
            summary_row(ws2, row, st, n); row += 1

    ws3 = wb.create_sheet("欠陥収束")
    write_convergence_sheet(ws3, bugs, "found_date", hdr_cell, border, to_serial)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"\n  ✅ エクスポート完了: {out}\n")


# ═══════════════════════════════════════════════════════════════════════════
# ISSUE コマンド
# ═══════════════════════════════════════════════════════════════════════════

def new_issue(issue_id):
    return {
        "id":            issue_id,
        "title":              None,
        "description":        None,
        "type":               None,
        "priority":           "中",
        "category":           None,
        "reporter":           None,
        "created_date":       date.today().isoformat(),
        "stage":              None,
        "planned_stage":      None,
        "status":             "未対応",
        "assignee":           None,
        "resolution":         None,
        "notes":              None,
        "resolved_date":      None,
        "verified_date":      None,
        "created_session_id": current_session_id(),
        "history":            [],
    }

def print_issue_row(issue):
    sid   = colored(issue["id"], BOLD)
    sstat = colored(f"{issue['status']:<5}", STATUS_COLOR.get(issue["status"], ""))
    spri  = colored(f"{issue['priority']:<3}", PRIORITY_COLOR.get(issue["priority"], ""))
    stype = (issue.get("type") or "-")[:6].ljust(6)
    title = issue["title"][:38].ljust(38)
    asgn  = (issue.get("assignee") or "-")[:8].ljust(8)
    print(f"  {sid}  {sstat}  {spri}  {stype}  {title}  {asgn}")

def issue_list(args):
    data   = load(ISSUE_DATA)
    issues = data["issues"]
    if args.status:   issues = [i for i in issues if i["status"]          == args.status]
    if args.priority: issues = [i for i in issues if i["priority"]        == args.priority]
    if args.type:     issues = [i for i in issues if i.get("type")        == args.type]
    if args.stage:    issues = [i for i in issues if i.get("stage")       == args.stage]
    if args.assignee: issues = [i for i in issues if i.get("assignee")    == args.assignee]
    if args.category: issues = [i for i in issues if i.get("category")    == args.category]
    if args.open:     issues = [i for i in issues if i["status"] not in ("解決済", "却下")]
    issues = sorted(issues, key=lambda i: ({"高":0,"中":1,"低":2}.get(i["priority"], 9), i["id"]))
    if not issues:
        print("該当する課題はありません。"); return
    print(f"\n  {BOLD}{'ID':<11}  {'状態':<7}  {'優先':<5}  {'種別':<8}  {'タイトル':<38}  {'担当者':<8}{RESET}")
    print("  " + "-" * 86)
    for i in issues:
        print_issue_row(i)
    print(f"\n  {DIM}合計: {len(issues)} 件{RESET}\n")

def issue_show(args):
    data  = load(ISSUE_DATA)
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    print(f"\n{'='*55}")
    print(f"  {BOLD}{issue['id']}  {issue['title']}{RESET}")
    print(f"{'='*55}")
    print(f"\n  {BOLD}説明:{RESET}\n    {issue.get('description') or '-'}")
    print(f"\n  種別          : {issue.get('type') or '-'}")
    print(f"  優先度        : {colored(issue['priority'], PRIORITY_COLOR.get(issue['priority'], ''))}")
    print(f"  カテゴリ      : {issue.get('category') or '-'}")
    print(f"  報告者        : {issue.get('reporter') or '-'}")
    print(f"  作成日        : {issue.get('created_date') or '-'}")
    print(f"  発生ステージ  : {issue.get('stage') or '-'}")
    print(f"  対応予定ステージ: {issue.get('planned_stage') or '-'}")
    print(f"  ステータス    : {colored(issue['status'], STATUS_COLOR.get(issue['status'], ''))}")
    print(f"  担当者        : {issue.get('assignee') or '-'}")
    if issue.get("resolution"):
        print(f"\n  {BOLD}対応方針:{RESET}\n    {issue['resolution']}")
    if issue.get("notes"):
        print(f"\n  {BOLD}備考:{RESET}\n    {issue['notes']}")
    print(f"\n  対応完了日    : {issue.get('resolved_date') or '-'}")
    print(f"  解決確認日    : {issue.get('verified_date') or '-'}")
    csid = issue.get("created_session_id") or "-"
    print(f"  登録セッション: {csid[:8] if csid != '-' else '-'}")
    history = issue.get("history", [])
    if history:
        print(f"\n  {BOLD}変更履歴:{RESET}")
        for h in history:
            sid = h.get("session_id") or ""
            sid_str = f"  [{sid[:8]}]" if sid else ""
            print(f"    {DIM}{h['date'][:10]}{RESET}  {h['actor']}  {h['from']} → {h['to']}{sid_str}")
            if h.get("comment"):
                print(f"             {DIM}{h['comment']}{RESET}")
    print()

def issue_add(args):
    data     = load(ISSUE_DATA)
    issue_id = next_id(data, "ISSUE")
    if args.title:
        issue = new_issue(issue_id)
        issue.update({
            "title":       args.title,
            "status":      args.status   or "未対応",
            "priority":    args.priority or "中",
            "type":        args.type,
            "assignee":    args.assignee,
            "reporter":    args.reporter,
            "stage":         args.stage,
            "planned_stage": args.planned_stage,
            "category":      args.category,
            "description":   args.description,
            "resolution":    args.resolution,
            "notes":         args.notes,
        })
        data["issues"].append(issue)
        save(ISSUE_DATA, data)
        print(f"\n  ✅ {colored(issue_id, BOLD)} を登録しました: {args.title}\n")
        return
    print(f"\n{BOLD}新規課題を登録します{RESET}\n")
    print(f"  ID: {colored(issue_id, BOLD)}\n")
    issue = new_issue(issue_id)
    issue.update({
        "title":       ask_interactive("タイトル",
                           guide="主題を一行で。例: SRS: RTM の追加 / merge.py: 削除候補判定の改修"),
        "description": ask_interactive("説明",
                           required=False,
                           guide="課題（問い）の背景・現状・困りごとを記述。例: 現在 XX の動作が YY になっており…"),
        "type":        ask_interactive("種別",
                           choices=VALID_TYPES, required=False,
                           guide="改善=既存仕様の変更/整理、調査=方針を決めるための情報収集、設計=実装方針・アーキの判断"),
        "priority":    ask_interactive("優先度",
                           choices=VALID_PRIORITIES, default="中",
                           guide="高=今すぐ対処が必要、中=近い将来対処、低=余裕のある時に対処"),
        "category":    ask_interactive("カテゴリ",
                           choices=VALID_CATEGORIES, required=False,
                           guide="どのモジュール・機能に関係するか"),
        "reporter":    ask_interactive("報告者",
                           required=False,
                           guide="人名 または モデル名（Opus / Sonnet 等）。Claude が登録する場合はモデル名"),
        "stage":       ask_interactive("発生ステージ",
                           choices=VALID_STAGES, required=False,
                           guide="課題が発生したフェーズ。URD=要求/SRS=要件/HLD=基本設計/LLD=詳細設計/COD=実装"),
        "status":      ask_interactive("ステータス",
                           choices=VALID_STATUSES, default="未対応"),
        "assignee":    ask_interactive("担当者",       required=False),
        "resolution":  ask_interactive("対応方針",
                           required=False,
                           guide="仕様検討の決着・対応方針。まだ未定なら空欄でも可"),
        "notes":       ask_interactive("備考",         required=False),
    })
    data["issues"].append(issue)
    save(ISSUE_DATA, data)
    print(f"\n  ✅ {colored(issue_id, BOLD)} を登録しました\n")

def issue_update(args):
    data  = load(ISSUE_DATA)
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    changed = []
    if args.status and args.status != issue["status"]:
        append_history(issue, args.actor, issue["status"], args.status, args.comment)
        old = issue["status"]
        issue["status"] = args.status
        changed.append(f"ステータス: {old} → {args.status}")
        if args.status == "対応完了" and not issue.get("resolved_date"):
            issue["resolved_date"] = date.today().isoformat()
        if args.status == "却下" and not issue.get("verified_date"):
            issue["verified_date"] = date.today().isoformat()
    for field, val in [
        ("priority",    args.priority),
        ("type",        args.type),
        ("stage",         args.stage),
        ("planned_stage", args.planned_stage),
        ("assignee",      args.assignee),
        ("category",    args.category),
        ("description", args.description),
        ("resolution",  args.resolution),
        ("notes",       args.notes),
    ]:
        if val is not None:
            changed.append(f"{field}: {issue.get(field)} → {val}")
            issue[field] = val
    if not changed:
        print("変更はありませんでした。"); return
    save(ISSUE_DATA, data)
    print(f"\n  ✅ {args.id} を更新しました:")
    for c in changed:
        print(f"     • {c}")
    print()

def issue_close(args):
    class _A: pass
    a = _A()
    a.id = args.id; a.status = "対応完了"; a.actor = args.actor; a.comment = args.comment
    a.priority = a.type = a.stage = a.planned_stage = a.assignee = a.category = None
    a.description = a.resolution = a.notes = None
    issue_update(a)

def issue_verify(args):
    data  = load(ISSUE_DATA)
    issue = next((i for i in data["issues"] if i["id"] == args.id), None)
    if not issue:
        print(f"エラー: {args.id} が見つかりません"); sys.exit(1)
    today = date.today().isoformat()
    issue["verified_date"] = today
    append_history(issue, args.actor, issue["status"], "解決済", f"[解決確認] {args.comment or ''}".strip())
    issue["status"] = "解決済"
    save(ISSUE_DATA, data)
    print(f"\n  ✅ {args.id} を解決済にしました（解決確認日: {today}）\n")

def issue_summary(args):
    data   = load(ISSUE_DATA)
    issues = data["issues"]
    total  = len(issues)
    print(f"\n  {BOLD}{'─'*40}\n  課題管理サマリー  (合計 {total} 件)\n  {'─'*40}{RESET}")
    print(f"\n  {BOLD}ステータス別{RESET}")
    for s in VALID_STATUSES:
        n = sum(1 for i in issues if i["status"] == s)
        print(f"  {colored(s.ljust(6), STATUS_COLOR.get(s,''))}  {colored(('█'*n).ljust(20), STATUS_COLOR.get(s,''))}  {n}")
    open_issues = [i for i in issues if i["status"] not in ("解決済", "却下")]
    print(f"\n  {BOLD}優先度別 (未解決){RESET}")
    for pv in VALID_PRIORITIES:
        n = sum(1 for i in open_issues if i["priority"] == pv)
        print(f"  {colored(pv.ljust(3), PRIORITY_COLOR.get(pv,''))}  {colored(('█'*n).ljust(20), PRIORITY_COLOR.get(pv,''))}  {n}")
    print(f"\n  {BOLD}種別別 (未解決){RESET}")
    for t in VALID_TYPES:
        n = sum(1 for i in open_issues if i.get("type") == t)
        bar = "█" * n
        print(f"  {t.ljust(6)}  {bar.ljust(20)}  {n}")
    print(f"\n  {BOLD}ステージ別 (未解決){RESET}")
    for st in VALID_STAGES:
        n = sum(1 for i in open_issues if i.get("stage") == st)
        if n:
            print(f"  {st.ljust(5)}  {'█'*n}  {n}")
    resolved = sum(1 for i in issues if i["status"] == "解決済")
    print(f"\n  {BOLD}KPI{RESET}")
    print(f"  未解決件数  : {len(open_issues)}")
    print(f"  優先度「高」: {colored(str(sum(1 for i in open_issues if i['priority']=='高')), PRIORITY_COLOR['高'])}")
    print(f"  解決率      : {resolved}/{total} ({resolved/total*100:.1f}%)" if total else "  解決率      : 0/0 (-)")
    print()

def issue_export(args):
    out = args.output or ISSUE_REPORT
    if args.if_changed and os.path.exists(out):
        if os.path.getmtime(ISSUE_DATA) <= os.path.getmtime(out):
            print(f"  変更なし、{out} の再生成をスキップ")
            return

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxlが必要です: pip install openpyxl"); sys.exit(1)

    data   = load(ISSUE_DATA)
    issues = data["issues"]
    wb     = Workbook()
    border, hdr_cell, data_cell, section_header, summary_row, make_dv, to_serial = make_excel_helpers()

    ws1 = wb.active
    ws1.title = "課題一覧"
    columns = [
        ("ID",           "id",            12),
        ("タイトル",     "title",         40),
        ("説明",         "description",   44),
        ("種別",         "type",          12),
        ("優先度",       "priority",      10),
        ("カテゴリ",     "category",      16),
        ("報告者",       "reporter",      14),
        ("作成日",       "created_date",  14),
        ("発生ステージ",   "stage",          16),
        ("対応予定ステージ", "planned_stage", 16),
        ("ステータス",   "status",        12),
        ("担当者",       "assignee",      14),
        ("対応方針",     "resolution",    30),
        ("備考",         "notes",         30),
        ("対応完了日",   "resolved_date", 14),
        ("解決確認日",   "verified_date", 14),
    ]
    ws1.row_dimensions[1].height = 28
    for i, (label, _, width) in enumerate(columns, 1):
        hdr_cell(ws1, 1, i, label, width=width)
    priority_bg = {"高": "FFE0E0", "中": "FFFBE0", "低": "E8F5E9"}
    pri_col = next(i for i, (_, f, _) in enumerate(columns, 1) if f == "priority")
    for ri, issue in enumerate(issues, 2):
        bg = "F0F0F0" if issue["status"] in ("解決済", "却下") else ("EBF3FB" if ri % 2 == 0 else "FFFFFF")
        for ci, (_, field, _) in enumerate(columns, 1):
            val = issue.get(field) or ""
            cell_bg = priority_bg.get(val, bg) if ci == pri_col else bg
            data_cell(ws1, ri, ci, val, bg=cell_bg, bold=(ci == pri_col and val in priority_bg))
        ws1.row_dimensions[ri].height = 40
    ws1.freeze_panes = "B2"
    last_row = len(issues) + 1
    for choices, field in [(VALID_STATUSES, "status"), (VALID_PRIORITIES, "priority"),
                           (VALID_TYPES, "type"), (VALID_STAGES, "stage"), (VALID_STAGES, "planned_stage")]:
        col_idx = next(i for i, (_, f, _) in enumerate(columns, 1) if f == field)
        ws1.add_data_validation(make_dv(choices, get_column_letter(col_idx), last_row))

    ws2 = wb.create_sheet("サマリー")
    open_issues = [i for i in issues if i["status"] not in ("解決済", "却下")]
    resolved    = sum(1 for i in issues if i["status"] == "解決済")
    total       = len(issues)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 12
    row = 1
    section_header(ws2, row, 1, "KPI"); row += 1
    summary_row(ws2, row, "合計件数",   total);           row += 1
    summary_row(ws2, row, "未解決件数", len(open_issues)); row += 1
    summary_row(ws2, row, "解決済件数", resolved);        row += 1
    summary_row(ws2, row, "解決率", f"{resolved/total*100:.1f}%" if total else "-%"); row += 2
    section_header(ws2, row, 1, "ステータス別"); row += 1
    for s in VALID_STATUSES:
        summary_row(ws2, row, s, sum(1 for i in issues if i["status"] == s)); row += 1
    row += 1
    section_header(ws2, row, 1, "優先度別（未解決）"); row += 1
    for pv in VALID_PRIORITIES:
        summary_row(ws2, row, pv, sum(1 for i in open_issues if i["priority"] == pv)); row += 1
    row += 1
    section_header(ws2, row, 1, "種別別（未解決）"); row += 1
    for t in VALID_TYPES:
        n = sum(1 for i in open_issues if i.get("type") == t)
        if n:
            summary_row(ws2, row, t, n); row += 1
    row += 1
    section_header(ws2, row, 1, "カテゴリ別（未解決）"); row += 1
    for cat in VALID_CATEGORIES:
        n = sum(1 for i in open_issues if i.get("category") == cat)
        if n:
            summary_row(ws2, row, cat, n); row += 1
    row += 1
    section_header(ws2, row, 1, "ステージ別（未解決）"); row += 1
    for st in VALID_STAGES:
        n = sum(1 for i in open_issues if i.get("stage") == st)
        if n:
            summary_row(ws2, row, st, n); row += 1

    ws3 = wb.create_sheet("収束グラフ")
    write_convergence_sheet(ws3, issues, "created_date", hdr_cell, border, to_serial)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"\n  ✅ エクスポート完了: {out}\n")


# ═══════════════════════════════════════════════════════════════════════════
# main
# ═══════════════════════════════════════════════════════════════════════════

def build_bug_parser(sub):
    p = sub.add_parser("bug", help="バグ管理")
    s = p.add_subparsers(dest="cmd")

    pl = s.add_parser("list");     pl.add_argument("--status"); pl.add_argument("--severity"); pl.add_argument("--assignee"); pl.add_argument("--category"); pl.add_argument("--open", action="store_true")
    ps = s.add_parser("show");     ps.add_argument("id")
    pa = s.add_parser("add");      pa.add_argument("--title"); pa.add_argument("--status", choices=VALID_STATUSES); pa.add_argument("--severity", choices=VALID_SEVERITIES); pa.add_argument("--found_stage", choices=VALID_STAGES); pa.add_argument("--stage", choices=VALID_STAGES); pa.add_argument("--category", choices=VALID_CATEGORIES); pa.add_argument("--found_in"); pa.add_argument("--cause"); pa.add_argument("--assignee"); pa.add_argument("--reporter"); pa.add_argument("--repro"); pa.add_argument("--description"); pa.add_argument("--resolution"); pa.add_argument("--notes")
    pu = s.add_parser("update");   pu.add_argument("id"); pu.add_argument("--status", choices=VALID_STATUSES); pu.add_argument("--severity", choices=VALID_SEVERITIES); pu.add_argument("--found_stage", choices=VALID_STAGES); pu.add_argument("--stage", choices=VALID_STAGES); pu.add_argument("--category", choices=VALID_CATEGORIES); pu.add_argument("--found_in"); pu.add_argument("--cause"); pu.add_argument("--repro"); pu.add_argument("--assignee"); pu.add_argument("--description"); pu.add_argument("--resolution"); pu.add_argument("--notes"); pu.add_argument("--actor", default="不明"); pu.add_argument("--comment")
    pc = s.add_parser("close");    pc.add_argument("id"); pc.add_argument("--actor", default="不明"); pc.add_argument("--comment")
    pv = s.add_parser("verify");   pv.add_argument("id"); pv.add_argument("--actor", default="不明"); pv.add_argument("--comment")
    s.add_parser("summary")
    pe = s.add_parser("export");   pe.add_argument("--output"); pe.add_argument("--if-changed", action="store_true")

    p.set_defaults(dispatch={
        "list": bug_list, "show": bug_show, "add": bug_add,
        "update": bug_update, "close": bug_close, "verify": bug_verify,
        "summary": bug_summary, "export": bug_export,
    }, parser=p)

def build_issue_parser(sub):
    p = sub.add_parser("issue", help="課題管理")
    s = p.add_subparsers(dest="cmd")

    pl = s.add_parser("list");     pl.add_argument("--status"); pl.add_argument("--priority"); pl.add_argument("--type"); pl.add_argument("--stage"); pl.add_argument("--assignee"); pl.add_argument("--category"); pl.add_argument("--open", action="store_true")
    ps = s.add_parser("show");     ps.add_argument("id")
    pa = s.add_parser("add");      pa.add_argument("--title"); pa.add_argument("--status", choices=VALID_STATUSES); pa.add_argument("--priority", choices=VALID_PRIORITIES); pa.add_argument("--type", choices=VALID_TYPES); pa.add_argument("--stage", choices=VALID_STAGES); pa.add_argument("--planned_stage", choices=VALID_STAGES); pa.add_argument("--category", choices=VALID_CATEGORIES); pa.add_argument("--assignee"); pa.add_argument("--reporter"); pa.add_argument("--description"); pa.add_argument("--resolution"); pa.add_argument("--notes")
    pu = s.add_parser("update");   pu.add_argument("id"); pu.add_argument("--status", choices=VALID_STATUSES); pu.add_argument("--priority", choices=VALID_PRIORITIES); pu.add_argument("--type", choices=VALID_TYPES); pu.add_argument("--stage", choices=VALID_STAGES); pu.add_argument("--planned_stage", choices=VALID_STAGES); pu.add_argument("--category", choices=VALID_CATEGORIES); pu.add_argument("--assignee"); pu.add_argument("--description"); pu.add_argument("--resolution"); pu.add_argument("--notes"); pu.add_argument("--actor", default="不明"); pu.add_argument("--comment")
    pc = s.add_parser("close");    pc.add_argument("id"); pc.add_argument("--actor", default="不明"); pc.add_argument("--comment")
    pv = s.add_parser("verify");   pv.add_argument("id"); pv.add_argument("--actor", default="不明"); pv.add_argument("--comment")
    s.add_parser("summary")
    pe = s.add_parser("export");   pe.add_argument("--output"); pe.add_argument("--if-changed", action="store_true")

    p.set_defaults(dispatch={
        "list": issue_list, "show": issue_show, "add": issue_add,
        "update": issue_update, "close": issue_close, "verify": issue_verify,
        "summary": issue_summary, "export": issue_export,
    }, parser=p)

def main():
    p   = argparse.ArgumentParser(description="バグ・課題 統合管理CLI")
    sub = p.add_subparsers(dest="kind")
    build_bug_parser(sub)
    build_issue_parser(sub)

    args = p.parse_args()
    if not args.kind:
        p.print_help(); return

    if not hasattr(args, "cmd") or args.cmd is None:
        args.parser.print_help(); return

    fn = args.dispatch.get(args.cmd)
    if fn:
        if getattr(args, "id", None) is not None:
            args.id = normalize_id(args.id, args.kind)
        fn(args)
    else:
        args.parser.print_help()

if __name__ == "__main__":
    main()
